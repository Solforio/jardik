"""`jardik forecast run|runs|review` on a synthetic extraction."""

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from fake_extraction import write_fake_extraction
from jardik.cli import app
from jardik.core.paths import JardikPaths
from synthetic import make_merged_frame

runner = CliRunner()

CONFIG = """\
name: cli_gbm
extraction_id: offices-test
split: {train_end: 2024-03-15, val_end: 2024-04-10, seed: 7}
lightgbm: {n_estimators: 100, num_leaves: 15, min_child_samples: 20, early_stopping_rounds: 15}
"""


@pytest.fixture
def project(tmp_path, monkeypatch):
    root = tmp_path / "project"
    (root / "pyproject.toml").parent.mkdir(parents=True)
    (root / "pyproject.toml").write_text("[project]\nname='x'\n")
    monkeypatch.setenv("JARDIK_HOME", str(root))
    merged = make_merged_frame(n_buildings=16, start="2024-01-01", days=150, seed=11)
    write_fake_extraction(JardikPaths.from_root(root), merged)
    (root / "run.yaml").write_text(CONFIG, encoding="utf-8")
    return root


def _run(project, *extra):
    return runner.invoke(app, ["forecast", "run", "--config", str(project / "run.yaml"), *extra])


def test_the_forecast_group_is_listed():
    result = runner.invoke(app, ["forecast", "--help"])
    assert result.exit_code == 0
    for command in ("run", "runs", "review"):
        assert command in result.stdout


def test_run_prints_the_verdict_and_a_metrics_table(project):
    result = _run(project)
    assert result.exit_code == 0, result.output
    assert "PASSED" in result.stdout and "lightgbm" in result.stdout
    assert "seasonal_naive_168" in result.stdout and "MAE" in result.stdout
    assert (project / "outputs" / "runs").is_dir()


def test_rerun_needs_force_and_runs_are_listed(project):
    assert _run(project).exit_code == 0
    again = _run(project)
    assert again.exit_code == 1 and "force" in again.output
    assert _run(project, "--force").exit_code == 0
    listing = runner.invoke(app, ["forecast", "runs"])
    assert "cli_gbm-" in listing.stdout and "PASSED" in listing.stdout


def test_invalid_config_and_missing_extraction(project, tmp_path):
    bad = project / "bad.yaml"
    bad.write_text(
        "name: x\nextraction_id: e\nsplit: {train_end: 2024-05-01, val_end: 2024-04-01}\n"
    )
    assert runner.invoke(app, ["forecast", "run", "--config", str(bad)]).exit_code == 2
    missing = project / "missing.yaml"
    missing.write_text(CONFIG.replace("offices-test", "nope-123"), encoding="utf-8")
    result = runner.invoke(app, ["forecast", "run", "--config", str(missing)])
    assert result.exit_code == 1 and "dataset extractions" in result.output


def test_review_builds_a_readable_folder(project):
    assert _run(project).exit_code == 0
    result = runner.invoke(app, ["forecast", "review"])
    assert result.exit_code == 0, result.output
    folders = list((project / "review").iterdir())
    assert len(folders) == 1 and folders[0].name.startswith("cli_gbm-")
    names = {p.name for p in folders[0].iterdir()}
    assert {"forecast_report.pdf", "forecast_tables.xlsx", "figures", "json"} <= names

    book = load_workbook(folders[0] / "forecast_tables.xlsx")
    for sheet in (
        "Read me",
        "Metrics",
        "Macro by model",
        "Skill",
        "Per building",
        "Feature importance",
        "Split roles",
        "Example week",
        "Limitations",
    ):
        assert sheet in book.sheetnames, sheet
    metrics = book["Metrics"]
    assert metrics["A1"].font.bold and metrics.freeze_panes == "A2"
    readme = "\n".join(str(c.value) for row in book["Read me"].iter_rows() for c in row if c.value)
    assert "PASSED" in readme and "CC BY-SA 4.0" in readme


def test_review_without_runs_or_with_unknown_id(project):
    assert runner.invoke(app, ["forecast", "review"]).exit_code == 1
    assert _run(project).exit_code == 0
    assert runner.invoke(app, ["forecast", "review", "nope-1"]).exit_code == 1
