"""`jardik dataset review`: a readable copy of an extraction (Excel workbook, PDF, sample CSV)."""

from datetime import UTC, date, datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from jardik.core.paths import JardikPaths
from jardik.dataset_manager.adapters.bdg2 import Bdg2Adapter, Bdg2Config
from jardik.dataset_manager.export.review import build_review
from jardik.dataset_manager.pipeline import run_extraction
from jardik.dataset_manager.request import ExtractionRequest


@pytest.fixture
def extraction(mini_bdg2: Path, tmp_path: Path):
    adapter = Bdg2Adapter(Bdg2Config(data_dir=mini_bdg2, usage_map={"Office": "office"}))
    paths = JardikPaths.from_root(tmp_path / "project")
    request = ExtractionRequest(
        name="rev",
        start=date(2016, 3, 8),
        end=date(2016, 3, 20),
        usage_categories=["office"],
        area={"site_ids": ["East"]},
        n_buildings=3,
        seed=1,
    )
    return run_extraction(request, adapter, paths, now=lambda: datetime(2026, 10, 16, tzinfo=UTC))


def test_review_folder_contains_pdf_workbook_and_sample(extraction, tmp_path):
    dest = build_review(extraction.output_dir, tmp_path / "review")
    names = {p.name for p in dest.iterdir()}
    assert {"extraction_report.pdf", "extraction_tables.xlsx", "merged_sample.csv"} <= names
    assert (dest / "extraction_report.pdf").read_bytes()[:4] == b"%PDF"
    assert dest.name == extraction.extraction_id


def test_workbook_sheets_are_readable_tables(extraction, tmp_path):
    dest = build_review(extraction.output_dir, tmp_path / "review")
    book = load_workbook(dest / "extraction_tables.xlsx")
    assert book.sheetnames[0] == "Read me"
    for sheet in (
        "Funnel",
        "Selected buildings",
        "Quality by building",
        "Weather variables",
        "Sites",
        "Hourly profile",
        "Merged sample",
    ):
        assert sheet in book.sheetnames, sheet
    selected = book["Selected buildings"]
    assert selected.max_row == 1 + 3  # header + 3 buildings
    assert selected["A1"].font.bold and selected.freeze_panes == "A2"
    readme = "\n".join(str(c.value) for row in book["Read me"].iter_rows() for c in row if c.value)
    assert extraction.extraction_id in readme and "CC BY-SA 4.0" in readme


def test_merged_sample_is_one_building_around_a_dst_change(extraction, tmp_path):
    dest = build_review(extraction.output_dir, tmp_path / "review")
    sample = pd.read_csv(dest / "merged_sample.csv", sep=";", decimal=",")
    assert sample["building_id"].nunique() == 1
    days = pd.to_datetime(sample["ts_local"]).dt.date
    assert date(2016, 3, 13) in set(days)  # US daylight saving started that day
    day = sample[days == date(2016, 3, 13)]
    assert len(day) == 23  # a 23 hour local day, shown as it is
    assert sample["target_kwh"].dtype.kind == "f"


def test_review_can_be_rebuilt_in_place(extraction, tmp_path):
    first = build_review(extraction.output_dir, tmp_path / "review")
    second = build_review(extraction.output_dir, tmp_path / "review")
    assert first == second and (second / "extraction_tables.xlsx").is_file()
