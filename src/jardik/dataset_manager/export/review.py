"""A readable copy of an extraction: Excel workbook, the PDF, a sample CSV and the JSON reports.

The extraction folder holds machine-friendly files (Parquet, large CSV). This builds something to
look at: one workbook with a sheet per table, the PDF report, and a small sample of the merged
training dataset (one building around a daylight-saving change) in Spanish CSV format.
"""

from __future__ import annotations

import json
import shutil
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from jardik.dataset_manager.export.writer import PLAIN_CSV, SPANISH_CSV

_HEADER_FILL = PatternFill("solid", fgColor="E8EEF5")
_MAX_ROWS = 5000
_SAMPLE_DAYS = 14


def _read_table(directory: Path, stem: str, manifest: dict[str, Any]) -> pd.DataFrame:
    export = manifest["export"]
    if export["format"] == "parquet":
        return pd.read_parquet(directory / f"{stem}.parquet")
    dialect = SPANISH_CSV if export.get("csv_style") == "spanish" else PLAIN_CSV
    frame = pd.read_csv(directory / f"{stem}.csv", sep=dialect["sep"], decimal=dialect["decimal"])
    for column in ("ts_utc", "issued_at_utc"):
        if column in frame.columns:
            frame[column] = pd.to_datetime(frame[column], utc=True)
    if "ts_local" in frame.columns:
        frame["ts_local"] = pd.to_datetime(frame["ts_local"])
    return frame


def _excel_safe(frame: pd.DataFrame) -> pd.DataFrame:
    """Excel has no timezone-aware datetimes: show UTC instants as plain UTC datetimes."""
    out = frame.copy()
    for column in out.columns:
        if isinstance(out[column].dtype, pd.DatetimeTZDtype):
            out[column] = out[column].dt.tz_convert("UTC").dt.tz_localize(None)
    return out


def _style(sheet: Any, frame: pd.DataFrame) -> None:
    sheet.freeze_panes = "A2"
    for index, column in enumerate(frame.columns, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        longest = max([len(str(column))] + [len(str(v)) for v in frame[column].head(200).tolist()])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 8), 60)
        if "ratio" in str(column):
            for row in range(2, len(frame) + 2):
                sheet.cell(row=row, column=index).number_format = "0.0%"
        elif pd.api.types.is_float_dtype(frame[column]):
            for row in range(2, len(frame) + 2):
                sheet.cell(row=row, column=index).number_format = "#,##0.00"


def _sample_window(
    merged: pd.DataFrame, quality: dict[str, Any], building: str
) -> tuple[pd.Timestamp, pd.Timestamp]:
    """Local-time window of the sample: around the building's first DST change, else the start."""
    days = [d for d in quality["weather_matching"]["dst_days"] if d["building_id"] == building]
    rows = merged[merged["building_id"] == building]
    if days:
        centre = pd.Timestamp(days[0]["local_date"])
        return centre - pd.Timedelta(days=7), centre + pd.Timedelta(days=7)
    start = rows["ts_local"].min().normalize()
    return start, start + pd.Timedelta(days=_SAMPLE_DAYS)


def _readme_rows(manifest: dict[str, Any], sheets: list[tuple[str, str]]) -> pd.DataFrame:
    request, dataset, counts = manifest["request"], manifest["dataset"], manifest["counts"]
    rows: list[tuple[str, str]] = [
        ("Extraction", manifest["extraction_id"]),
        ("Created", manifest["created_at"]),
        ("Dataset", dataset["name"]),
        ("Licence", dataset["licence"]),
        ("Period (UTC)", f"{request['start']} to {request['end']} (end excluded)"),
        ("Buildings selected", str(counts["buildings"])),
        ("Hourly rows", str(counts["load_rows"])),
        ("Attribution", manifest["attribution"]),
        ("", ""),
        ("Sheets", ""),
        *sheets,
        ("", ""),
        (
            "About numbers",
            "Cells hold real numbers; Excel shows them with your regional settings "
            "(decimal comma in a Spanish locale). Ratios are shown as percentages.",
        ),
        (
            "About times",
            "ts_utc is UTC. ts_local is the local wall-clock time of the building "
            "(daylight saving included) and utc_offset_min is its offset.",
        ),
    ]
    return pd.DataFrame(rows, columns=["Item", "Value"])


def build_review(extraction_dir: Path, dest_root: Path) -> Path:
    """Create `<dest_root>/<extraction_id>/` with the readable files and return it."""
    manifest = json.loads((extraction_dir / "extraction_manifest.json").read_text(encoding="utf-8"))
    quality = json.loads((extraction_dir / "quality_report.json").read_text(encoding="utf-8"))
    distribution = json.loads(
        (extraction_dir / "distribution_report.json").read_text(encoding="utf-8")
    )
    dest = dest_root / manifest["extraction_id"]
    dest.mkdir(parents=True, exist_ok=True)

    table = lambda stem: _read_table(extraction_dir, stem, manifest)  # noqa: E731
    merged = table("merged_training_dataset")
    funnel = pd.DataFrame(manifest["funnel"])
    excluded = table("excluded_buildings")

    building = manifest["selected_buildings"][0]["building_id"]
    low, high = _sample_window(merged, quality, building)
    in_window = (merged["ts_local"] >= low) & (merged["ts_local"] < high)
    sample = merged[(merged["building_id"] == building) & in_window].reset_index(drop=True)

    missing = quality["weather_matching"]["stats"]["value_missing_ratio"]
    rules = manifest["weather"]["value_rules"]
    weather_vars = pd.DataFrame(
        [
            {
                "variable": v,
                "unit": manifest["weather"]["units"][v],
                "missing_ratio": missing.get(v),
                "value_rules_applied": "; ".join(f"{k}: {c}" for k, c in rules.get(v, {}).items())
                or "none",
            }
            for v in manifest["weather"]["variables"]
        ]
    )
    sites = pd.DataFrame(manifest["geographic_coverage"])
    profile = pd.DataFrame(
        {
            "hour_local": list(map(int, distribution["hourly_profile_kwh"])),
            "mean_kwh": list(distribution["hourly_profile_kwh"].values()),
        }
    )

    sheets: list[tuple[str, pd.DataFrame, str]] = [
        ("Funnel", funnel, "How many buildings each filter kept and removed."),
        (
            "Selected buildings",
            table("selected_buildings"),
            "The buildings of this extraction, with gap statistics.",
        ),
        ("Metadata", table("building_metadata"), "Static description of the selected buildings."),
        (
            "Excluded buildings",
            excluded.head(_MAX_ROWS),
            f"Buildings dropped by quality or weather (first {_MAX_ROWS}).",
        ),
        (
            "Quality by building",
            pd.DataFrame(quality["buildings"]),
            "Missing, imputed and gap statistics.",
        ),
        ("Sites", sites, "Geographic coverage by site."),
        ("Weather variables", weather_vars, "Units, missing share and value rules."),
        ("Hourly profile", profile, "Mean load by local hour of day."),
        ("Columns", pd.DataFrame(manifest["columns"]), "Documentation of merged_training_dataset."),
        ("Merged sample", sample, f"One building ({building}) around a daylight-saving change."),
    ]

    workbook_path = dest / "extraction_tables.xlsx"
    guide = [(name, description) for name, _, description in sheets]
    readme = _readme_rows(manifest, guide)
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="Read me", index=False)
        sheet = writer.book["Read me"]
        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 110
        for row in sheet.iter_rows(min_row=2):
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
            row[0].font = Font(bold=True)
        _style_header_only(sheet)
        for name, frame, _ in sheets:
            safe = _excel_safe(frame)
            safe.to_excel(writer, sheet_name=name, index=False)
            _style(writer.book[name], safe)

    sample.to_csv(dest / "merged_sample.csv", index=False, sep=";", decimal=",")
    for name in ("extraction_report.pdf",):
        if (extraction_dir / name).is_file():
            shutil.copy2(extraction_dir / name, dest / name)
    json_dir = dest / "json"
    json_dir.mkdir(exist_ok=True)
    for name in ("extraction_manifest.json", "quality_report.json", "distribution_report.json"):
        shutil.copy2(extraction_dir / name, json_dir / name)
    return dest


def _style_header_only(sheet: Any) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
